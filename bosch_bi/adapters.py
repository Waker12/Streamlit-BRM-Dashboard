"""Adapters for the actual 2026 workbooks inspected in this conversation.

Workbook text is data. Only the structural contracts below control execution.
No formulas, macros, URLs or instructions embedded in a workbook are executed.
"""
from dataclasses import dataclass,field
from io import BytesIO
import calendar
import re
import pandas as pd
import openpyxl
import xlrd

@dataclass
class Adapted:
    kind: str
    datasets: dict = field(default_factory=dict)
    notes: list = field(default_factory=list)
    checks: list = field(default_factory=list)

SHOPEE_HEADER=['Date','Sales (MYR)','Sales (Shopee Rebate applied)','Orders','Sales per Order','Product Clicks','Visitors','Order Conversion Rate','Cancelled Orders','Cancelled Sales','Returned/Refunded Orders','Returned/Refunded Sales','# of buyers','# of new buyers','# of existing buyers','# of potential buyers','Repeat Purchase Rate']
LAZADA_HEADER=['Date','Revenue','Visitors','Buyers','Orders','Pageviews','Units Sold','Conversion Rate','Revenue per Buyer','Visitor Value','Add to Cart Users','Add to Cart Units','Wishlists','Wishlist Users','Average Order Value','Average Basket Size','Cancelled Amount','Return/Refund Amount']

def num(value):
    if value is None:return None
    text=str(value).strip()
    if text in ['', '-', 'N/A']:return None
    cleaned=text.replace(',','').removesuffix('%')
    result=float(cleaned)
    if not pd.notna(result) or abs(result)==float('inf'):raise ValueError('Non-finite source number')
    return result

def interval(text,dayfirst=True):
    regex=r'\d{2}-\d{2}-\d{4}' if dayfirst else r'\d{4}-\d{2}-\d{2}'
    values=re.findall(regex,str(text))
    if len(values)!=2:raise ValueError('Expected an explicit two-date reporting interval in the source.')
    fmt='%d-%m-%Y' if dayfirst else '%Y-%m-%d'
    a,b=[pd.to_datetime(v,format=fmt).date() for v in values]
    if a.day!=1 or b.day!=calendar.monthrange(b.year,b.month)[1] or (a.year,a.month)!=(b.year,b.month):raise ValueError('This adapter expects a complete single calendar month.')
    return str(a),str(b)

def frames(result):
    result.datasets={k:pd.DataFrame(v) for k,v in result.datasets.items() if len(v)}
    return result

def adapter_kind(raw):
    if raw[:8]==bytes.fromhex('D0CF11E0A1B11AE1'):
        w=xlrd.open_workbook(file_contents=raw,on_demand=True)
        return 'lazada' if 'Key Metrics' in w.sheet_names() else ('lazada_products' if 'Product' in w.sheet_names() else None)
    if raw[:2]==b'PK':
        import zipfile
        with zipfile.ZipFile(BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('Expanded workbook exceeds 100 MB.')
        w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True)
        if all(n in w.sheetnames for n in ['Monthly Tracking','BICOS','Sales & Traffic Shopee','Sales & Traffic Shopee Lazada','Ads 2026']):return 'brm'
        if 'Top Performing Products' in w.sheetnames:return 'shopee_products'
        if all(n in w.sheetnames for n in ['Placed Order','Confirmed Order','Paid Order']):return 'shopee'
        if 'MY' in w.sheetnames and str(w['MY']['D11'].value).endswith('Actual') and str(w['MY']['D14'].value).endswith('Official Forecast'):return 'internal'
    return None

def parse_source(raw,kind=None,country='MY',period=None):
    kind=kind or adapter_kind(raw)
    if kind=='brm':
        from .brm import performance_workbook
        return performance_workbook(raw)
    if kind=='shopee_products':return shopee_products(raw,country,period)
    if kind=='lazada_products':return lazada_products(raw,country)
    if kind=='internal':return internal(raw)
    if kind=='shopee':return shopee(raw,country)
    if kind=='lazada':return lazada(raw,country)
    raise ValueError('No inspected source adapter matches this workbook. Use canonical mapping or review the new schema.')

def internal(raw):
    w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True)
    s=w['MY']
    if s['A3'].value!='MY':raise ValueError('The inspected internal adapter expects the MY country worksheet.')
    year=int(re.search(r'\d{4}',str(s['D11'].value))[0])
    if s['D12'].value!=f'{year-1} Actual' or s['D13'].value!=f'{year} Stretch Forecast' or s['D14'].value!=f'{year} Official Forecast':raise ValueError('Internal row labels changed; review the template.')
    if [s.cell(10,c).value for c in range(5,17)]!=[calendar.month_abbr[i] for i in range(1,13)]:raise ValueError('Internal month headers changed.')
    if any(s.cell(9,c).value!='EUR' for c in range(5,17)):raise ValueError('Internal currency headers changed.')
    r=Adapted('Bosch internal MY TNS',{'financial_actuals':[],'forecasts':[]})
    for month in range(1,13):
        col=month+4
        for row,y in [(11,year),(12,year-1)]:
            value=num(s.cell(row,col).value)
            if value is not None:r.datasets['financial_actuals'].append(dict(country='MY',marketplace='Consolidated',period=f'{y}-{month:02d}',currency='EUR',tns=value))
        for row,scenario in [(13,'stretch'),(14,'official')]:
            value=num(s.cell(row,col).value)
            if value is None:raise ValueError('A forecast formula lacks a cached numeric result. Recalculate and save the source in Excel first.')
            r.datasets['forecasts'].append(dict(country='MY',marketplace='Consolidated',period_start=f'{year}-{month:02d}-01',period_end=f'{year}-{month:02d}-{calendar.monthrange(year,month)[1]}',currency='EUR',scenario=scenario,tns=value))
    target=num(s['C3'].value)
    if target is None:raise ValueError('Missing internal annual target in C3 (EUR millions).')
    r.datasets['forecasts'].append(dict(country='MY',marketplace='Consolidated',period_start=f'{year}-01-01',period_end=f'{year}-12-31',currency='EUR',scenario='target',tns=target*1_000_000))
    official=sum(x['tns'] for x in r.datasets['forecasts'] if x['scenario']=='official')
    if abs(official-target*1_000_000)>.02:raise ValueError('Monthly official forecasts do not reconcile with the explicit annual target.')
    for row in [11,12,13,14]:
        source_total=num(s.cell(row,17).value)
        calculated=sum(num(s.cell(row,c).value) or 0 for c in range(5,17))
        r.checks.append(dict(metric=str(s.cell(row,4).value),source=source_total,calculated=calculated,status='Matches' if source_total is not None and abs(calculated-source_total)<.02 else 'Mismatch / missing cached total'))
    r.notes=['Country-level MY finance is stored as Consolidated; never allocated to Shopee or Lazada.','Only TNS actuals, TNS forecasts and the annual TNS target are present. GM, finance orders, return-rate and concession definitions are absent.','Cached Excel formula results are read; formulas are not executed. Future actual blanks remain unavailable.']
    return frames(r)

def shopee(raw,country):
    if country!='MY':raise ValueError('This inspected Shopee schema uses Sales (MYR); it cannot be assigned to Singapore.')
    w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True)
    r=Adapted('Shopee MY shop statistics',{'marketplace':[],'traffic':[],'marketplace_daily':[],'contributions':[]})
    report_periods=set()
    for basis in ['Placed Order','Confirmed Order','Paid Order']:
        sheet=w[basis];data=list(sheet.values)
        if list(data[0])!=SHOPEE_HEADER or list(data[3])!=SHOPEE_HEADER:raise ValueError(f'{basis}: source headers changed.')
        a,b=interval(data[1][0]);period=a[:7];report_periods.add(period)
        row=dict(zip(SHOPEE_HEADER,data[1]));scope=dict(country='MY',marketplace='Shopee',sales_basis=basis)
        r.datasets['marketplace'].append(dict(**scope,period=period,currency='MYR',gmv=num(row['Sales (MYR)']),orders=num(row['Orders']),reported_aov=num(row['Sales per Order'])))
        r.datasets['traffic'].append(dict(**scope,period_start=a,period_end=b,visitors=num(row['Visitors']),buyers=num(row['# of buyers']),reported_order_cr_pct=num(row['Order Conversion Rate']),product_clicks=num(row['Product Clicks'])))
        daily=[]
        for values in data[4:]:
            if not values[0]:continue
            d=pd.to_datetime(values[0],format='%d-%m-%Y').strftime('%Y-%m-%d')
            if not a<=d<=b:raise ValueError('Daily record outside declared report interval.')
            source=dict(zip(SHOPEE_HEADER,values))
            daily.append(dict(**scope,date=d,currency='MYR',gmv=num(source['Sales (MYR)']),orders=num(source['Orders']),visitors=num(source['Visitors']),buyers=num(source['# of buyers']),reported_order_cr_pct=num(source['Order Conversion Rate'])))
        if len(daily)!=pd.Timestamp(b).day or len({x['date'] for x in daily})!=len(daily):raise ValueError('Daily coverage is incomplete or duplicated.')
        for metric,source in [('gmv','Sales (MYR)'),('orders','Orders')]:
            calculated=sum(x[metric] for x in daily);expected=num(row[source])
            r.checks.append(dict(period=period,basis=basis,metric=metric,source=expected,calculated=calculated,status='Matches' if expected is not None and abs(calculated-expected)<.05 else 'Mismatch'))
        r.datasets['marketplace_daily']+=daily
        # Product contribution sheets are truncated excerpts per channel, not the complete store assortment.
        candidates=[n for n in w.sheetnames if n.startswith('Product Contribution') and w[n].cell(2,2).value==basis]
        if len(candidates)!=1:raise ValueError('Expected exactly one source product-contribution sheet per order basis.')
        channel=None
        for values in w[candidates[0]].values:
            if values[0] in ['Product Card','Seller Live','Seller Video','Shopee Affiliate']:channel=values[0];continue
            if channel and values[0] and str(values[0]).isdigit():
                r.datasets['contributions'].append(dict(**scope,period=period,currency='MYR',channel=channel,listing_sku=str(values[0]),product_name=values[1],revenue=num(values[4]),units=num(values[8]),source_orders=num(values[7]),buyers=num(values[12]),product_clicks=num(values[6]),unique_product_clicks=num(values[14]),reported_order_cr_pct=num(values[10])))
    if len(report_periods)!=1:raise ValueError('Order-basis sheets cover different months.')
    r.notes=['Each order basis is stored independently. Placed Order is the default; do not add Placed, Confirmed and Paid totals together.','Monthly summary audiences are used directly; daily audiences are stored separately and never summed into monthly unique totals.','Shopee Order Conversion Rate uses orders / product clicks, which differs from buyer / visitor CR. Both retain their distinct definitions.','Overall units sold, full SKU performance, add-to-cart units, approved SKU mapping and ads spend are not supplied by this export.','Channel product excerpts are retained separately and are not promoted to full SKU performance or cross-channel totals.']
    return frames(r)

def lazada(raw,country):
    if country!='MY':raise ValueError('The supplied Lazada workbooks are MY exports. Review a Singapore schema before using SGD.')
    w=xlrd.open_workbook(file_contents=raw)
    s=w.sheet_by_name('Key Metrics');data=[s.row_values(i) for i in range(s.nrows)]
    if list(data[5])!=LAZADA_HEADER:raise ValueError('Lazada Key Metrics row 6 headers changed.')
    a,b=interval(data[0][0]);sa,sb=interval(data[6][0],dayfirst=False)
    if (a,b)!=(sa,sb):raise ValueError('Lazada report header and summary intervals disagree.')
    scope=dict(country='MY',marketplace='Lazada',sales_basis='Lazada Revenue')
    summary=dict(zip(LAZADA_HEADER,data[6]));period=a[:7]
    r=Adapted('Lazada MY Business Advisor',{'marketplace':[],'traffic':[],'marketplace_daily':[]})
    r.datasets['marketplace'].append(dict(**scope,period=period,currency='MYR',gmv=num(summary['Revenue']),orders=num(summary['Orders']),units=num(summary['Units Sold']),reported_aov=num(summary['Average Order Value'])))
    r.datasets['traffic'].append(dict(**scope,period_start=a,period_end=b,visitors=num(summary['Visitors']),buyers=num(summary['Buyers']),reported_cr_pct=num(summary['Conversion Rate'])))
    daily=[]
    for values in data[8:]:
        d=pd.to_datetime(values[0],format='%Y-%m-%d').strftime('%Y-%m-%d')
        if not a<=d<=b:raise ValueError('Lazada daily row lies outside report interval.')
        row=dict(zip(LAZADA_HEADER,values))
        daily.append(dict(**scope,date=d,currency='MYR',gmv=num(row['Revenue']),orders=num(row['Orders']),units=num(row['Units Sold']),visitors=num(row['Visitors']),buyers=num(row['Buyers']),reported_cr_pct=num(row['Conversion Rate'])))
    if len(daily)!=pd.Timestamp(b).day or len({x['date'] for x in daily})!=len(daily):raise ValueError('Lazada daily coverage is incomplete or duplicated.')
    r.datasets['marketplace_daily']=daily
    for metric,source in [('gmv','Revenue'),('orders','Orders'),('units','Units Sold')]:
        calculated=sum(x[metric] for x in daily);expected=num(summary[source])
        r.checks.append(dict(period=period,basis='Lazada Revenue',metric=metric,source=expected,calculated=calculated,status='Matches' if expected is not None and abs(calculated-expected)<.05 else 'Mismatch'))
    r.notes=['Monthly summary is authoritative for unique audience; daily audiences are never added into monthly or YTD uniques.','The supplied Metric Definition sheet defines revenue as buyer paid amount plus platform discounts and shipping subsidies. It is marketplace GMV, not finance TNS.','This workbook contains store metrics and daily records, but no SKU performance, SKU master, ad spend or advertising-attributed revenue.']
    return frames(r)


LAZADA_PRODUCT_HEADER=['Product ID','Product Name','URL','Seller SKU','SKU ID','Product Visitors','Product Pageviews','Visitor Value','Add to Cart Users','Add to Cart Units','Add to Cart Conversion Rate','Wishlist Users','Wishlists','Buyers','Orders','Units Sold','Revenue','Conversion Rate','Revenue per Buyer','Revenue share']
SHOPEE_PRODUCT_HEADER=['Item ID','Product','Current Item Status','Variation ID','Variation Name','Current Variation Status','SKU','Parent SKU','Sales (Placed Order) (MYR)','Sales (Confirmed Order) (MYR)','Product Impression','Product Clicks','CTR','Order Conversion Rate (Placed Order)','Order Conversion Rate (Confirmed Order)','Placed Order','Confirmed Order','Units (Placed Order)','Units (Confirmed Order)','Buyers (Placed Order)','Buyers (Confirmed Order)','Conversion Rate (Placed Order)','Conversion Rate (Confirmed Order)','Sales per Order (Placed Order) (MYR)','Sales per Order (Confirmed Order) (MYR)','Unique Product Impressions','Unique Product Clicks','Product Visitors (Visit)','Product Page Views','Product Bounce Visitors','Product Bounce Rate','Search Clicks','Likes','Product Visitors (Add to Cart)','Units (Add to Cart)','Conversion Rate (Add to Cart)','Repeat Order Rate (Placed Order)','Repeat Order Rate (Confirmed Order)','Average Days to Repeat Order (Placed Order)','Average Days to Repeat Order (Confirmed Order)']

def period_from_filename(filename):
    """Suggestion only: user-supplied filenames are not embedded report metadata."""
    year=re.search(r'(?<!\d)(20\d{2})(?!\d)',filename)
    months={calendar.month_abbr[i].lower():i for i in range(1,13)}
    tokens=re.findall(r'[A-Za-z]+',filename.lower())
    month=next((months[t[:3]] for t in tokens if len(t)>=3 and t[:3] in months),None)
    return f'{year[0]}-{month:02d}' if year and month else None

def text_value(value):
    return None if value is None or str(value).strip() in ['', '-'] else str(value).strip()

def shopee_products(raw,country,period):
    if country!='MY':raise ValueError('Inspected Shopee product columns are explicitly MYR.')
    if not period or not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',period):raise ValueError('Shopee product files contain no report dates. Supply and verify the reporting month (YYYY-MM).')
    w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True)
    rows=list(w['Top Performing Products'].values)
    if list(rows[0])!=SHOPEE_PRODUCT_HEADER:raise ValueError('Shopee product header changed; review schema.')
    result=Adapted('Shopee MY product performance',{'listing_products':[],'listing_variants':[]})
    for values in rows[1:]:
        if not values[0]:continue
        row=dict(zip(SHOPEE_PRODUCT_HEADER,values));item=text_value(row['Item ID']);variant=text_value(row['Variation ID'])
        if item is None:raise ValueError('Missing Shopee item identity.')
        for basis in ['Placed Order','Confirmed Order']:
            base=dict(country='MY',marketplace='Shopee',period=period,currency='MYR',sales_basis=basis,listing_sku=item,product_name=row['Product'],revenue=num(row[f'Sales ({basis}) (MYR)']),units=num(row[f'Units ({basis})']),orders=num(row[basis]),buyers=num(row[f'Buyers ({basis})']),period_source='Explicit supplied month; file has no embedded report interval')
            if variant is None:
                base.update(visitors=num(row['Product Visitors (Visit)']),atc_units=num(row['Units (Add to Cart)']),reported_cr_pct=num(row[f'Conversion Rate ({basis})']),source_material=text_value(row['Parent SKU']))
                result.datasets['listing_products'].append(base)
            else:
                # Audience / ATC columns on variation rows are usually missing; retain that absence.
                base.update(variant_id=variant,variant_name=text_value(row['Variation Name']),seller_sku=text_value(row['SKU']),visitors=num(row['Product Visitors (Visit)']),atc_units=num(row['Units (Add to Cart)']))
                result.datasets['listing_variants'].append(base)
    result.notes=['Reporting month was explicitly supplied; this workbook contains no embedded reporting interval. Filename suggestions must be verified.','Only Top Performing Products is imported. Newly Listed, price and ads-recommendation sheets overlap and are not additional sales or ad-spend datasets.','Item-total rows and variation rows are stored separately; adding them would double count sales.','Seller / parent SKU text is retained as source metadata. Bosch canonical identity, BU, product group, bundle allocation and GM require an approved master / finance source.','Paid Order product metrics are absent; only Placed and Confirmed product bases are available.']
    return frames(result)

def lazada_products(raw,country):
    if country!='MY':raise ValueError('This inspected product adapter is for the supplied MY Lazada exports.')
    w=xlrd.open_workbook(file_contents=raw);s=w.sheet_by_name('Product')
    if s.row_values(5)!=LAZADA_PRODUCT_HEADER:raise ValueError('Lazada product headers changed; review schema.')
    start,end=interval(s.cell_value(0,0),dayfirst=False);period=start[:7]
    result=Adapted('Lazada MY product performance',{'listing_products':[],'listing_variants':[]})
    material_candidates={}
    for i in range(6,s.nrows):
        row=dict(zip(LAZADA_PRODUCT_HEADER,s.row_values(i)))
        item=text_value(row['Product ID']);variant=text_value(row['SKU ID'])
        if item is None:raise ValueError('Missing Lazada Product ID.')
        base=dict(country='MY',marketplace='Lazada',period=period,currency='MYR',sales_basis='Lazada Revenue',listing_sku=item,product_name=row['Product Name'],revenue=num(row['Revenue']),units=num(row['Units Sold']),orders=num(row['Orders']),visitors=num(row['Product Visitors']),buyers=num(row['Buyers']),atc_units=num(row['Add to Cart Units']),period_source='Embedded workbook reporting interval')
        if variant is None:
            base.update(reported_cr_pct=num(row['Conversion Rate']),source_material=None)
            result.datasets['listing_products'].append(base)
        else:
            seller=text_value(row['Seller SKU'])
            base.update(variant_id=variant,seller_sku=seller)
            result.datasets['listing_variants'].append(base)
            if seller:material_candidates.setdefault(item,set()).add(seller)
    for row in result.datasets['listing_products']:
        materials=material_candidates.get(row['listing_sku'],set())
        row['source_material']=next(iter(materials)) if len(materials)==1 else None
    result.notes=['The embedded workbook date interval is used as the reporting month.','Rows with SKU ID = dash are product totals; actual SKU ID rows are variants. They are stored separately to prevent double counting.','Variant-level product visitors and conversion rates are missing in this export. Product audiences are not allocated to individual variants.','Seller SKU is source metadata, retained without promoting it to an approved Bosch canonical mapping or BU.']
    return frames(result)
