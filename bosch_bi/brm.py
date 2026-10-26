"""Inspected five-tab BRM dependency adapter. Other workbook tabs are ignored."""
from io import BytesIO
import calendar
import openpyxl
from .adapters import Adapted,frames,num

MAIN_TABS=['Monthly Tracking','BICOS','Sales & Traffic Shopee','Sales & Traffic Shopee Lazada','Ads 2026']

def performance_workbook(raw):
    w=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True)
    formula=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=False)
    if not all(n in w.sheetnames for n in MAIN_TABS):raise ValueError('The five inspected BRM tabs are required.')
    r=Adapted('Bosch MY five-tab BRM', {k:[] for k in ['brm_tracking','financial_rates','brm_review','marketplace','traffic','ads','kpi_benchmarks']})
    tracking=w['Monthly Tracking']
    if tracking['A35'].value!=2025 or tracking['A52'].value!=2026:raise ValueError('Monthly Tracking year blocks changed.')
    for year,first in [(2025,37),(2026,54)]:
        for month in range(1,13):
            row=first+month-1;period=f'{year}-{month:02d}'
            if tracking.cell(row,1).value!=calendar.month_abbr[month]:raise ValueError('Monthly Tracking month headers changed.')
            for platform,cols in [('Lazada',dict(gmv=2,visitors=3,buyers=4,orders=5,units=7,reported_cr_pct=8,cancelled_amount=17,returned_amount=18)),('Shopee',dict(gmv=21,visitors=25,buyers=31,orders=22,reported_cr_pct=26,cancelled_amount=28,returned_amount=30))]:
                gmv=num(tracking.cell(row,cols['gmv']).value)
                if gmv is None:continue
                values={key:num(tracking.cell(row,col).value) for key,col in cols.items()}
                if values['reported_cr_pct'] is not None:values['reported_cr_pct']*=100
                basis='Lazada Revenue' if platform=='Lazada' else ('BRM historical · status unverified' if year==2025 else 'BRM mixed order basis · not primary')
                base=dict(country='MY',marketplace=platform,period=period,currency='MYR',sales_basis=basis)
                r.datasets['brm_tracking'].append(dict(**base,**values))
                # Only 2025 Lazada history has an identified comparable source definition.
                if year==2025 and platform=='Lazada':
                    r.datasets['marketplace'].append(dict(**base,gmv=gmv,orders=values['orders'],units=values['units']))
                    r.datasets['traffic'].append(dict(country='MY',marketplace=platform,sales_basis=basis,period_start=period+'-01',period_end=f'{period}-{calendar.monthrange(year,month)[1]}',visitors=values['visitors'],buyers=values['buyers'],reported_cr_pct=values['reported_cr_pct']))
    b=w['BICOS']
    if b['A23'].value!=2026 or b['A28'].value!='GM %':raise ValueError('BICOS 2026 GM row moved.')
    for month in range(1,13):
        gm=num(b.cell(28,month+1).value)
        if gm is not None:
            cell=b.cell(28,month+1).coordinate
            r.datasets['financial_rates'].append(dict(country='MY',marketplace='Consolidated',period=f'2026-{month:02d}',gm_pct=gm,source_cell='BICOS!'+cell))
        tns=num(b.cell(26,month+1).value)
        if tns is not None:r.datasets['brm_review'].append(dict(tab='BICOS',cell=b.cell(26,month+1).coordinate,period=f'2026-{month:02d}',metric='BICOS TNS reference',value=tns,detail='Reference only. Internal Bosch TNS workbook remains primary; reconcile differences.'))
    latest=max(int(x['period'][-2:]) for x in r.datasets['financial_rates'])
    for scenario,columns,end in [('target',[21,22,23,25,27],'2026-12-31'),('official',[31,32,33,35,37],f'2026-{latest:02d}-{calendar.monthrange(2026,latest)[1]}')]:
        values={key:num(b.cell(17,col).value) for key,col in zip(['gm_pct','visitors','buyers','orders','return_rate'],columns)}
        if values['return_rate'] is not None:values['return_rate']*=100
        r.datasets['kpi_benchmarks'].append(dict(country='MY',marketplace='Consolidated',period_start='2026-01-01',period_end=end,scenario=scenario,**values))
    r.datasets['brm_review'].append(dict(tab='BICOS',cell='N28',metric='Source-reported YTD GM %',value=num(b['N28'].value),detail='Workbook uses '+str(formula['BICOS']['N28'].value)+'. Not an approved weighted YTD rate; monthly bases are needed.'))
    r.datasets['brm_review'].append(dict(tab='BICOS',cell='A15',metric='Historical BICOS excluded',detail='User explicitly excludes 2025 and earlier BICOS data. Use internal workbook for 2025 TNS; Monthly Tracking for marketplace history.'))
    for platform,first in [('Lazada',3),('Shopee',21)]:
        for month in range(1,13):
            row=first+month-1;s=w['Ads 2026'];period=f'2026-{month:02d}'
            if s.cell(row,1).value!=calendar.month_abbr[month].upper():raise ValueError('Ads month rows moved.')
            spend=num(s.cell(row,2).value);gmv=num(s.cell(row,3).value)
            if spend is None and gmv is None:continue
            if spend is None or gmv is None:raise ValueError('Ads spend and GMV must both be available.')
            r.datasets['ads'].append(dict(country='MY',marketplace=platform,period=period,currency='MYR',spend=spend,ads_gmv=gmv,attribution_model=f'BRM reported · unverified attribution · {platform}'))
            cell=f'E{row}';ref=formula['Ads 2026'][cell].value
            r.datasets['brm_review'].append(dict(tab='Ads 2026',cell=cell,period=period,metric=platform+' ads store-GMV denominator',value=num(s[cell].value),detail=f'Source formula {ref} points to 2025 Monthly Tracking. Not used for 2026 spend/store GMV or ad share. Ad attribution/order status remains unverified.'))
    r.datasets['brm_review'].append(dict(tab='Monthly Tracking',cell='BC66',metric='Value-based return-rate definition',value=num(w['Monthly Tracking']['BC66'].value)*100,detail='Formula is combined Returned/Refunded Sales / combined store GMV. Legacy mixed-order inputs are reference only; not substituted for a finance rate or Confirmed-only metric.'))
    for tab in ['Sales & Traffic Shopee','Sales & Traffic Shopee Lazada']:
        r.datasets['brm_review'].append(dict(tab=tab,cell='C3',metric='Chart dependency',detail='2025 history comes from Monthly Tracking per user instruction, not hardcoded chart values. 2026 shop performance comes from raw exports; Shopee Confirmed Orders only.'))
    r.notes=['Only the five user-identified tabs are used; other tabs are ignored.','2025 BICOS and older series are excluded entirely. Internal Bosch file is authoritative for 2025 TNS.','2025 Monthly Tracking marketplace history is preserved. Shopee historical order status is unverified and is not mislabeled Confirmed.','2026 Monthly Tracking is retained for reconciliation only; raw seller-centre exports remain primary.','Monthly 2026 GM% is source-reported; no YTD weighting denominator is invented.','Ads spend and Ads GMV are entered source values; attribution is unverified. Wrong-year store-GMV formulas and combined advertising scopes are not reused.','Legacy workbook FX is not applied. Fixed Bosch EUR 1 = MYR 4.64 remains in force.']
    return frames(r)
