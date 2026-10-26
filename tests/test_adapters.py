from io import BytesIO
from pathlib import Path
import json
import os
import calendar
import pytest
import openpyxl
from bosch_bi.adapters import parse_source,adapter_kind,SHOPEE_HEADER,num
from bosch_bi.storage import Store,validate


def workbook_bytes(w):
    raw=BytesIO();w.save(raw);return raw.getvalue()

def shopee_fixture():
    w=openpyxl.Workbook();w.remove(w.active)
    for basis in ['Placed Order','Confirmed Order','Paid Order']:
        s=w.create_sheet(basis)
        s.append(SHOPEE_HEADER)
        s.append(['01-07-2026-31-07-2026','3,100.00','-',31,100,3100,500,'1.00%',0,0,0,0,20,20,0,30,'0%'])
        s.append([]);s.append(SHOPEE_HEADER)
        for day in range(1,32):s.append([f'{day:02d}-07-2026','100.00','-',1,100,100,50,'1.00%',0,0,0,0,1,1,0,1,'0%'])
        s=w.create_sheet('Product Contribution '+basis[:8])
        s.append(['Date','Order Type']);s.append(['01-07-2026-31-07-2026',basis]);s.append([])
        s.append(['Product Card']);s.append(['Item ID','Product'])
        s.append(['123','Synthetic drill','Normal','10%','310',100,10,'1.5',2,'10%','15%',200,2,50,8])
    return workbook_bytes(w)

def test_source_numeric_parser_distinguishes_missing():
    assert num('1,234.56')==1234.56
    assert num('0.00')==0
    assert num('-') is None
    assert num('0.70%')==.7
    with pytest.raises(ValueError):num('instructions are not code')

def test_shopee_preserves_order_bases_daily_grain_and_excerpt_limits():
    raw=shopee_fixture()
    assert adapter_kind(raw)=='shopee'
    result=parse_source(raw)
    assert len(result.datasets['marketplace'])==3
    assert len(result.datasets['marketplace_daily'])==93
    assert len(result.datasets['contributions'])==3
    assert set(result.datasets['marketplace'].sales_basis)=={'Placed Order','Confirmed Order','Paid Order'}
    assert set(result.datasets['traffic'].visitors)=={500}
    assert all(check['status']=='Matches' for check in result.checks)
    assert result.datasets['contributions'].source_orders.iloc[0]==1.5
    assert 'products' not in result.datasets and 'ads' not in result.datasets
    assert result.datasets['traffic'].reported_order_cr_pct.iloc[0]==1
    assert 'reported_cr_pct' not in result.datasets['traffic']
    for key,frame in result.datasets.items():assert not validate(key,frame)[1]

def test_shopee_header_drift_is_rejected():
    raw=shopee_fixture();w=openpyxl.load_workbook(BytesIO(raw));w['Placed Order']['B1']='Unknown Sales'
    with pytest.raises(ValueError,match='headers changed'):parse_source(workbook_bytes(w))

def test_internal_finance_is_country_consolidated_and_blank_future_is_missing():
    w=openpyxl.Workbook();s=w.active;s.title='MY';s['A3']='MY';s['C3']=.012
    for row,label in [(11,'2026 Actual'),(12,'2025 Actual'),(13,'2026 Stretch Forecast'),(14,'2026 Official Forecast')]:s.cell(row,4,label)
    for month in range(1,13):
        col=month+4;s.cell(9,col,'EUR');s.cell(10,col,calendar.month_abbr[month])
        if month<=9:s.cell(11,col,500)
        s.cell(12,col,400);s.cell(13,col,1100);s.cell(14,col,1000)
    for row,value in [(11,4500),(12,4800),(13,13200),(14,12000)]:s.cell(row,17,value)
    r=parse_source(workbook_bytes(w))
    assert set(r.datasets['financial_actuals'].marketplace)=={'Consolidated'}
    assert len(r.datasets['financial_actuals'])==21
    target=r.datasets['forecasts'].query("scenario=='target'")
    assert target.tns.iloc[0]==12000
    assert r.datasets['financial_actuals'].query("period=='2026-10'").empty
    assert 'orders' not in r.datasets['financial_actuals']

@pytest.mark.skipif(not os.environ.get('BOSCH_SOURCE_MANIFEST'),reason='Private supplied workbooks are not bundled with source')
def test_all_supplied_workbooks_and_idempotent_imports(tmp_path):
    files=json.loads(Path(os.environ['BOSCH_SOURCE_MANIFEST']).read_text())
    store=Store(tmp_path/'real-sources.sqlite')
    for file in files:
        source=parse_source(Path(file['path']).read_bytes())
        for key,df in source.datasets.items():assert not validate(key,df)[1],(file['key'],key)
        assert all(check['status']=='Matches' for check in source.checks),file['key']
        result=store.ingest_bundle(source.datasets,Path(file['path']).name)
        assert result['status']=='imported',(file['key'],result)
        assert store.ingest_bundle(source.datasets,'renamed.xls')['status']=='duplicate'
    assert len(store.read('marketplace'))==36
    assert len(store.read('traffic'))==36
    assert len(store.read('financial_actuals'))==21
    assert len(store.read('forecasts'))==25
    assert len(store.read('marketplace_daily'))==1092
    assert len(store.read('contributions'))==540
    assert store.read('products').empty
    assert store.read('ads').empty

def test_shopee_product_adapter_requires_explicit_period_and_separates_variants():
    from bosch_bi.adapters import SHOPEE_PRODUCT_HEADER
    w=openpyxl.Workbook();s=w.active;s.title='Top Performing Products'
    s.append(SHOPEE_PRODUCT_HEADER)
    values={k:'0' for k in SHOPEE_PRODUCT_HEADER}
    values.update({'Item ID':'123','Product':'Synthetic product','Variation ID':'-','Variation Name':'-','Parent SKU':'SOURCE-SKU','Sales (Placed Order) (MYR)':'200','Sales (Confirmed Order) (MYR)':'100','Product Visitors (Visit)':'100','Units (Placed Order)':'2','Units (Confirmed Order)':'1','Buyers (Placed Order)':'2','Buyers (Confirmed Order)':'1','Units (Add to Cart)':'10'})
    s.append([values[k] for k in SHOPEE_PRODUCT_HEADER])
    values.update({'Variation ID':'789','Variation Name':'Size A','SKU':'SOURCE-VARIANT','Product Visitors (Visit)':'-','Units (Add to Cart)':'-'})
    s.append([values[k] for k in SHOPEE_PRODUCT_HEADER])
    raw=workbook_bytes(w)
    assert adapter_kind(raw)=='shopee_products'
    with pytest.raises(ValueError,match='no report dates'):parse_source(raw)
    r=parse_source(raw,period='2026-09')
    assert len(r.datasets['listing_products'])==2
    assert len(r.datasets['listing_variants'])==2
    assert r.datasets['listing_products'].revenue.sum()==300
    assert r.datasets['listing_variants'].visitors.isna().all()
    assert 'products' not in r.datasets and 'sku_master' not in r.datasets

@pytest.mark.skipif(not os.environ.get('BOSCH_PRODUCT_MANIFEST'),reason='Private supplied product workbooks are not bundled')
def test_supplied_product_exports_do_not_double_count_parents_and_variants(tmp_path):
    from bosch_bi.adapters import period_from_filename
    from bosch_bi.engine import product_store_reconciliation
    files=json.loads(Path(os.environ['BOSCH_PRODUCT_MANIFEST']).read_text())
    store=Store(tmp_path/'products.sqlite')
    for f in files:
        result=parse_source(Path(f['path']).read_bytes(),period=period_from_filename(Path(f['path']).name))
        for key,df in result.datasets.items():assert not validate(key,df)[1],(f['key'],key)
        imported=store.ingest_bundle(result.datasets,Path(f['path']).name)
        assert imported['status']=='imported',(f['key'],imported)
        assert store.ingest_bundle(result.datasets,'renamed.xls')['status']=='duplicate'
    products=store.read('listing_products');variants=store.read('listing_variants')
    assert len(products)==19247
    assert len(variants)==3538
    assert not products.duplicated(['country','marketplace','period','sales_basis','listing_sku']).any()
    assert variants.loc[variants.marketplace=='Lazada','visitors'].isna().all()
    assert store.read('products').empty and store.read('sku_master').empty
    if os.environ.get('BOSCH_SOURCE_MANIFEST'):
        for f in json.loads(Path(os.environ['BOSCH_SOURCE_MANIFEST']).read_text()):
            result=parse_source(Path(f['path']).read_bytes())
            assert store.ingest_bundle(result.datasets,Path(f['path']).name)['status']=='imported'
        checks=product_store_reconciliation(products,store.read('marketplace'))
        flags=checks[checks.Status=='Mismatch']
        assert set(flags.Period)=={'2026-01'}
        assert set(flags.Marketplace)=={'Shopee'}
